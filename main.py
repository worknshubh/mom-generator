from dotenv import load_dotenv
load_dotenv()
from fastapi import FastAPI,File,UploadFile
from fastapi.responses import FileResponse
from docx import Document
from reportlab.pdfgen import canvas
app = FastAPI()
from langchain.chat_models import init_chat_model
import shutil
model = init_chat_model(model='mistral-small-latest')
from langchain_core.prompts import ChatPromptTemplate
from transformers import pipeline
pipe = pipeline("automatic-speech-recognition", model="openai/whisper-base")

transcription = ""
@app.post('/upload')
async def audio_upload(file:UploadFile= File(...)):
    temp_path = f"temp_{file.filename}"
    with open(temp_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
        
    result = pipe(temp_path,return_timestamps=True,language="en")
    transcription = result['text']
    minutes_of_the_meeting_prompt = ChatPromptTemplate.from_messages([
    ("system", """You are a professional meeting minutes assistant. Your task is to extract and summarize key information from meeting transcriptions.

    RULES:
    1. Only include information explicitly mentioned in the transcription
    2. Do NOT add assumptions, interpretations, or external knowledge
    3. Do NOT hallucinate or invent any details
    4. If something is unclear or ambiguous, note it as "[Unclear]"
    5. Organize information in clear bullet points
    6. Focus on actionable items, decisions, and key discussion points

    FORMAT YOUR RESPONSE AS FOLLOWS:

    **Meeting Summary**
    - Brief overview of the meeting's main purpose

    **Key Discussion Points**
    - Main topics discussed (only what was actually mentioned)
    - Important questions raised
    - Concerns or challenges identified

    **Decisions Made**
    - Specific decisions or conclusions reached
    - Include who made the decision if mentioned

    **Action Items**
    - Tasks to be completed
    - Responsible person (if mentioned)
    - Deadlines (if mentioned)

    **Next Steps**
    - Follow-up meetings or activities planned

    If any section has no relevant information from the transcription, write "None mentioned" for that section."""),
        
        ("user", """Please generate meeting minutes from the following transcription. Remember to only include information that is explicitly stated in the transcription.

    TRANSCRIPTION:
    {transcription}

    Generate the meeting minutes following the specified format.""")
    ])

    mom_final_prompt = minutes_of_the_meeting_prompt.invoke({
        'transcription':transcription
    })

    output = model.invoke(mom_final_prompt)
    print(output.content)
    
    doc = Document()
    doc.add_heading('Meeting Minutes', 0)

    for line in output.content.split('\n'):
        if line.strip():
            if line.startswith('**') and line.endswith('**'):
               
                doc.add_heading(line.strip('*'), level=1)
            else:
               
                doc.add_paragraph(line)

   
    doc_file = "meeting_minutes.docx"
    doc.save(doc_file)

    return FileResponse(doc_file, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", filename="meeting_minutes.docx")